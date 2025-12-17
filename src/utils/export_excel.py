import os
import pandas as pd
from datetime import datetime
from openpyxl.styles import Font
from src.utils.interface import console  # <--- Импортируем консоль Rich

def save_wallets_to_excel(wallets_data):
    if not wallets_data:
        console.print("[yellow]⚠️ Нет данных для сохранения в Excel.[/yellow]")
        return

    rows = []
    for w in wallets_data:
        address = w.get("wallet")
        pnl_val = w.get("realized_pnl")
        
        rows.append({
            "Wallet": address,
            "Total Predictions": w.get("total_predictions", 0),
            "Active Predictions": w.get("active_predictions", 0),
            "Realized PnL ($)": pnl_val if pnl_val is not None else 0,
            "Polymarket": "Open Profile",
            "Hashdive": "Analyze"
        })

    df = pd.DataFrame(rows)

    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    if not os.path.exists("results"):
        os.makedirs("results")
    
    filename = f"results/wallets_{timestamp}.xlsx"

    try:
        with pd.ExcelWriter(filename, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name="Wallets")
            
            workbook = writer.book
            worksheet = writer.sheets["Wallets"]
            link_font = Font(color="0000FF", underline="single")

            for i, w in enumerate(wallets_data):
                row_idx = i + 2
                address = w.get("wallet")
                
                poly_cell = worksheet.cell(row=row_idx, column=5)
                poly_cell.value = "Open Profile"
                poly_cell.hyperlink = f"https://polymarket.com/@{address}"
                poly_cell.font = link_font
                
                hash_cell = worksheet.cell(row=row_idx, column=6)
                hash_cell.value = "Analyze"
                hash_cell.hyperlink = f"https://hashdive.com/Analyze_User?user_address={address}"
                hash_cell.font = link_font

            for idx, col_name in enumerate(df.columns):
                col_letter = chr(65 + idx)
                if col_name == "Wallet": width = 45
                elif col_name in ["Polymarket", "Hashdive"]: width = 15
                else: width = 20
                worksheet.column_dimensions[col_letter].width = width
                
                if col_name == "Realized PnL ($)":
                    for r in range(2, len(df) + 2):
                        cell = worksheet.cell(row=r, column=idx+1)
                        cell.number_format = '"$"#,##0.00'

        # Используем console.print вместо print
        console.print(f"[green]💾 Таблица успешно сохранена:[/green] [bold]{filename}[/bold]")
        console.print(f"[dim]📊 Всего записей: {len(df)}[/dim]")
        
    except Exception as e:
        console.print(f"[red]❌ Ошибка при сохранении Excel: {e}[/red]")